"""Anti-colisão de férias direto na planilha de controle.

Mesmas regras do botão "Aplicar anti-colisão" do painel:
- limite de pessoas fora ao mesmo tempo por função (ou grupo), padrão 10% do quadro;
- limite conjunto de gestores fora ao mesmo tempo;
- mantém a duração; início de seg a qua, fora de feriado e dos 2 dias antes dele;
- respeita 12 meses de casa e o prazo concessivo; não mexe em período já iniciado
  nem em colaborador travado; procura a data mais próxima da original.

Gera uma CÓPIA da planilha trocando só INÍCIO/DIAS (F,G,I,J,L,M) das linhas alteradas.
Fórmulas, formatação, calendário e abas ficam intactos; o Excel recalcula ao abrir.

Uso:
  python -I scripts/anti_colisao.py [planilha.xlsx] [--regras ferias_regras.json] [--hoje AAAA-MM-DD] [--simular]
"""
import argparse
import csv
import datetime as dt
import json
import math
import re
import sys
import zipfile
from pathlib import Path

import openpyxl

EPOCH = dt.date(1970, 1, 1)
EXCEL_OFFSET = 25569  # serial do Excel em 01/01/1970
DOW = ['dom', 'seg', 'ter', 'qua', 'qui', 'sex', 'sáb']
REGRAS_PADRAO = {
    "limitar_por": "funcao",        # "funcao" ou "grupo"
    "max_gestores_juntos": 2,        # 0 desliga a regra conjunta da gestão
    "deslocamento_preferido_dias": 90,
    "limites": {},                   # ex.: {"TECNICO DE REDE EXTERNA": 5}
    "travados": [],                  # nomes que o anti-colisão não altera
}


def dn(d):
    return (d - EPOCH).days


def dd(n):
    return EPOCH + dt.timedelta(days=n)


def br(n):
    return dd(n).strftime('%d/%m/%Y')


def dow(n):
    return (n + 4) % 7  # 0 = domingo


def add_y(n, k):
    d = dd(n)
    try:
        return dn(d.replace(year=d.year + k))
    except ValueError:  # 29/02 vira 01/03, como no Date.UTC do painel
        return dn(dt.date(d.year + k, 3, 1))


def to_day(v):
    if isinstance(v, dt.datetime):
        return dn(v.date())
    if isinstance(v, dt.date):
        return dn(v)
    if isinstance(v, (int, float)) and v > 20000:
        return round(v) - EXCEL_OFFSET
    return None


def grupo(f):
    f = (f or '').upper()
    if re.search(r'GESTOR|GERENTE|ANALISTA|COORDENADOR|SUPERVISOR', f):
        return 'Gestão'
    if 'B2C' in f:
        return 'Técnico B2C'
    if 'B2B' in f:
        return 'Técnico B2B'
    if 'INFRAESTRUTURA' in f:
        return 'Infraestrutura'
    if 'REDE' in f:
        return 'Rede Externa'
    return 'Outros'


# ---------- leitura ----------
def ler_aba(ws):
    out = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i < 7:
            continue
        r = list(r) + [None] * 13
        if str(r[0] or '').startswith('Equipe em f'):
            break
        nome = r[1].strip() if isinstance(r[1], str) else ''
        adm = to_day(r[2])
        if not nome or adm is None:
            continue
        per = []
        for a, b in ((5, 6), (8, 9), (11, 12)):
            ini = to_day(r[a])
            dias = None if r[b] in (None, '') else int(r[b])
            if ini is not None or dias is not None:
                per.append((ini, dias))
        out.append({'nome': nome, 'adm': adm, 'funcao': str(r[3] or '').strip(), 'base': r[4] or '', 'p': per})
    return out


def ler_planilha(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = wb.sheetnames
    s_g = next((n for n in names if re.search('GEST', n, re.I)), None)
    s_e = next((n for n in names if re.search('f[ée]rias', n, re.I) and not re.search('GEST', n, re.I)), names[0])
    s_f = next((n for n in names if re.search('feriad', n, re.I)), None)
    fer = []
    if s_f:
        for i, r in enumerate(wb[s_f].iter_rows(values_only=True)):
            if i < 4:
                continue
            r = list(r) + [None] * 5
            d = to_day(r[0])
            if d is not None:
                fer.append({'data': d, 'nome': r[1], 'abr': str(r[2] or ''), 'cons': str(r[4] or '').upper()})
    equipe = ler_aba(wb[s_e])
    gestao = ler_aba(wb[s_g]) if s_g else []
    wb.close()
    return equipe, gestao, fer


# ---------- modelo ----------
def montar(equipe, gestao, fer, avisos):
    hol = [h for h in fer if h['cons'] == 'SIM']

    def hol_for(base):
        b = str(base or '').upper()
        return {h['data'] for h in hol if h['abr'] == 'Nacional' or (b and h['abr'].upper().endswith(' ' + b))}

    def dedupe(lst, aba):
        seen, out = {}, []
        for x in lst:
            if x['nome'] in seen:
                igual = seen[x['nome']] == x
                avisos.append(f"{x['nome']} aparece mais de uma vez na aba {aba}"
                              + (" com os mesmos dados (contado uma vez)" if igual else " com dados diferentes (usada a primeira linha)"))
                continue
            seen[x['nome']] = x
            out.append(x)
        return out

    eq, ge = dedupe(equipe, 'Equipe'), dedupe(gestao, 'GESTÃO')
    eq_map = {x['nome']: x for x in eq}
    for g in ge:
        e = eq_map.get(g['nome'])
        if e and e['p'] != g['p']:
            avisos.append(f"{g['nome']} tem períodos diferentes entre as abas; o ajuste usa a aba Equipe e grava o resultado nas duas")
    todos = eq + [g for g in ge if g['nome'] not in eq_map]
    pessoas = []
    for x in todos:
        per = [{'ini': i, 'dias': d} for i, d in x['p'] if i is not None and d]
        per.sort(key=lambda q: q['ini'])
        pessoas.append({**x, 'H': hol_for(x['base']), 'grupo': grupo(x['funcao']),
                        'aniv': add_y(x['adm'], 1), 'orig': [dict(q) for q in per], 'per': per})
    return pessoas


def conc_win(p, s):
    k = 0
    while add_y(p['adm'], k + 1) <= s:
        k += 1
    if k == 0:
        return None
    return add_y(p['adm'], k), add_y(p['adm'], k + 1) - 1


def bad_start(H, n):
    w = dow(n)
    return w >= 4 or w == 0 or n in H or (n + 1) in H or (n + 2) in H


def caps_for(pessoas, regras):
    modo = regras['limitar_por']
    key = (lambda p: 'F|' + p['funcao']) if modo == 'funcao' else (lambda p: 'G|' + p['grupo'])
    n = {}
    for p in pessoas:
        n[key(p)] = n.get(key(p), 0) + 1
    over = {k.upper(): v for k, v in regras.get('limites', {}).items()}
    caps = {k: int(over.get(k[2:].upper(), max(1, math.ceil(c * 0.10)))) for k, c in n.items()}
    gc = int(regras['max_gestores_juntos'])
    caps['GESTAO'] = gc

    def keys(p):
        ks = [key(p)]
        if gc > 0 and p['grupo'] == 'Gestão':
            ks.append('GESTAO')
        return ks
    return n, caps, keys


def medir(pessoas, caps, keys, r0, length):
    cnt = {}
    for p in pessoas:
        ks = keys(p)
        for q in p['per']:
            for d in range(q['ini'], q['ini'] + q['dias']):
                i = d - r0
                if 0 <= i < length:
                    for k in ks:
                        cnt.setdefault(k, [0] * length)[i] += 1
    exc, dias, pico = 0, set(), {}
    for k, a in cnt.items():
        pico[k] = max(a)
        for i, c in enumerate(a):
            if c > caps[k]:
                exc += c - caps[k]
                dias.add(i)
    return {'excesso': exc, 'dias': len(dias), 'pico': pico}


def anti_colisao(pessoas, regras, hoje):
    n, caps, keys = caps_for(pessoas, regras)
    travados = {t.strip().upper() for t in regras.get('travados', [])}
    lo = min([q['ini'] for p in pessoas for q in p['per']] + [dn(dt.date(2027, 1, 1))])
    r0 = dn(dd(lo).replace(day=1))
    length = dn(dt.date(2028, 12, 31)) - r0 + 800
    antes = medir(pessoas, caps, keys, r0, length)
    cnt = {}

    def arr(k):
        return cnt.setdefault(k, [0] * length)

    def add(ks, s, d):
        for x in range(s, s + d):
            i = x - r0
            if 0 <= i < length:
                for k in ks:
                    arr(k)[i] += 1

    def excess_of(ks, s, d):
        e = 0
        for x in range(s, s + d):
            i = x - r0
            if 0 <= i < length:
                for k in ks:
                    c = arr(k)[i]
                    if c >= caps[k]:
                        e += c - caps[k] + 1
        return e

    placed = {id(p): [] for p in pessoas}
    fixed, free = [], []
    for p in pessoas:
        lk = p['nome'].upper() in travados
        for q in (p['per'] if lk else p['orig']):
            cw = conc_win(p, q['ini']) or (p['aniv'], add_y(p['aniv'], 1) - 1)
            it = {'p': p, 'dias': q['dias'], 'o': q['ini'], 'ini': q['ini'], 'cw': cw}
            (fixed if lk or q['ini'] <= hoje else free).append(it)
    for it in fixed:
        add(keys(it['p']), it['ini'], it['dias'])
        placed[id(it['p'])].append(it)

    w0, w1 = dn(dt.date(2027, 1, 1)), dn(dt.date(2028, 1, 31))

    def valid(it, s):
        p, d = it['p'], it['dias']
        if s <= hoje or bad_start(p['H'], s) or s < p['aniv']:
            return False
        if it['cw'] and (s < it['cw'][0] or s + d - 1 > it['cw'][1]):
            return False
        if s < min(w0, it['o']) or s + d - 1 > max(w1, it['o'] + d - 1):
            return False
        return all(s + d - 1 < o['ini'] - 1 or s > o['ini'] + o['dias'] for o in placed[id(p)])

    free.sort(key=lambda it: (it['o'], -it['dias']))
    rej = []
    for it in free:
        ks = keys(it['p'])
        if valid(it, it['o']) and excess_of(ks, it['o'], it['dias']) == 0:
            add(ks, it['o'], it['dias'])
            placed[id(it['p'])].append(it)
        else:
            rej.append(it)
    rej.sort(key=lambda it: (it['dias'], it['o']))
    rej.reverse()  # igual ao painel: ordena e inverte (empates em ordem inversa)
    max_shift = int(regras['deslocamento_preferido_dias'])
    sem_vaga = []
    for it in rej:
        ks = keys(it['p'])
        best, bs = None, math.inf
        for s in range(it['o'] - 370, it['o'] + 371):
            sh = abs(s - it['o'])
            if sh >= bs or not valid(it, s):
                continue
            e = excess_of(ks, s, it['dias'])
            sc = (1e6 + e * 1e3 if e else 0) + sh + (1e5 if sh > max_shift else 0) + (0.5 if s < it['o'] else 0)
            if sc < bs:
                bs, best = sc, s
        if best is None:
            best = it['o']
            sem_vaga.append((it, 'nenhuma data válida'))
        elif bs >= 1e6:
            sem_vaga.append((it, 'sem data livre: ficou na de menor choque'))
        it['ini'] = best
        add(ks, best, it['dias'])
        placed[id(it['p'])].append(it)
    for p in pessoas:
        if p['nome'].upper() in travados:
            continue
        p['per'] = sorted(({'ini': it['ini'], 'dias': it['dias']} for it in placed[id(p)]), key=lambda q: q['ini'])
    depois = medir(pessoas, caps, keys, r0, length)
    return antes, depois, sem_vaga, n, caps


# ---------- escrita (cópia da planilha) ----------
def decode_xml(s):
    s = re.sub(r'&#x([0-9a-f]+);', lambda m: chr(int(m.group(1), 16)), s, flags=re.I)
    s = re.sub(r'&#(\d+);', lambda m: chr(int(m.group(1))), s)
    return s.replace('&lt;', '<').replace('&gt;', '>').replace('&quot;', '"').replace('&apos;', "'").replace('&amp;', '&')


def gravar(src, dst, alterados):
    zin = zipfile.ZipFile(src)
    wbx = zin.read('xl/workbook.xml').decode('utf-8')
    rels = zin.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    rel = {}
    for m in re.finditer(r'<Relationship\b[^>]*>', rels):
        i, t = re.search(r'\bId="([^"]+)"', m.group(0)), re.search(r'\bTarget="([^"]+)"', m.group(0))
        if i and t:
            rel[i.group(1)] = 'xl/' + re.sub(r'^/?xl/', '', t.group(1)).lstrip('/')
    sheets = [(decode_xml(re.search(r'\bname="([^"]+)"', m.group(0)).group(1)), re.search(r'\br:id="([^"]+)"', m.group(0)).group(1))
              for m in re.finditer(r'<sheet\b[^>]*>', wbx)]
    alvo = [rel[r] for nm, r in sheets if re.search('f[ée]rias', nm, re.I)]
    ss = []
    if 'xl/sharedStrings.xml' in zin.namelist():
        ssx = zin.read('xl/sharedStrings.xml').decode('utf-8')
        for m in re.finditer(r'<si>([\s\S]*?)</si>', ssx):
            body = re.sub(r'<rPh\b[\s\S]*?</rPh>', '', m.group(1))
            ss.append(decode_xml(''.join(re.findall(r'<t\b[^>]*>([\s\S]*?)</t>', body))))
    col = {'F': (0, 'i'), 'G': (0, 'd'), 'I': (1, 'i'), 'J': (1, 'd'), 'L': (2, 'i'), 'M': (2, 'd')}
    linhas, achados, faltando = 0, set(), []
    novos = {}
    for path in alvo:
        x = zin.read(path).decode('utf-8')
        row_p = {}
        for m in re.finditer(r'<c r="B(\d+)"([^>]*?)>(?:<f>[\s\S]*?</f>)?<v>([^<]*)</v></c>', x):
            if int(m.group(1)) < 8:
                continue
            nm = (ss[int(m.group(3))] if re.search(r'\bt="s"', m.group(2)) else decode_xml(m.group(3))).strip()
            if nm in alterados:
                row_p[m.group(1)] = alterados[nm]
                achados.add(nm)
        feitos, cel = set(), {}

        def troca(m):
            c, r, a, body = m.group(1), m.group(2), m.group(3), m.group(4)
            p = row_p.get(r)
            if not p or '<f' in body:
                return m.group(0)
            feitos.add(r)
            cel[(r, c)] = True
            j, t = col[c]
            q = p['per'][j] if j < len(p['per']) else None
            at = re.sub(r'\s+t="[^"]*"', '', a)
            if not q:
                return f'<c r="{c}{r}"{at}/>'
            return f'<c r="{c}{r}"{at}><v>{q["ini"] + EXCEL_OFFSET if t == "i" else q["dias"]}</v></c>'
        x = re.sub(r'<c r="([FGIJLM])(\d+)"([^>]*?)(/>|>[\s\S]*?</c>)', troca, x)
        for r, p in row_p.items():
            for c, (j, t) in col.items():
                if j < len(p['per']) and (r, c) not in cel:
                    faltando.append(f'{p["nome"]}: célula {c}{r} não existe ou tem fórmula')
        linhas += len(feitos)
        novos[path] = x
    novos['xl/workbook.xml'] = wbx if 'fullCalcOnLoad' in wbx else re.sub(r'<calcPr\b([^>]*?)\s*/>', r'<calcPr\1 fullCalcOnLoad="1"/>', wbx)
    with zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = novos[info.filename].encode('utf-8') if info.filename in novos else zin.read(info.filename)
            zout.writestr(info, data, compress_type=zipfile.ZIP_DEFLATED)
    zin.close()
    nao_achados = sorted(set(alterados) - achados)
    return linhas, nao_achados, faltando


# ---------- principal ----------
def main():
    ap = argparse.ArgumentParser(description='Anti-colisão de férias na planilha de controle')
    ap.add_argument('planilha', nargs='?', default='Controle_Ferias_N2_2027.xlsx')
    ap.add_argument('--regras', default='ferias_regras.json')
    ap.add_argument('--hoje', help='data de referência AAAA-MM-DD (padrão: hoje)')
    ap.add_argument('--saida', default='relatorios')
    ap.add_argument('--simular', action='store_true', help='só calcula e gera o relatório, sem criar a planilha ajustada')
    a = ap.parse_args()

    src = Path(a.planilha)
    if not src.exists():
        sys.exit(f'Planilha não encontrada: {src}')
    rp = Path(a.regras)
    if not rp.exists():
        rp.write_text(json.dumps(REGRAS_PADRAO, ensure_ascii=False, indent=2), encoding='utf-8')
    regras = {**REGRAS_PADRAO, **json.loads(rp.read_text(encoding='utf-8'))}
    hoje = dn(dt.date.fromisoformat(a.hoje)) if a.hoje else dn(dt.date.today())

    avisos = []
    equipe, gestao, fer = ler_planilha(src)
    if not equipe:
        sys.exit('Nenhum colaborador encontrado a partir da linha 8 da aba de férias.')
    pessoas = montar(equipe, gestao, fer, avisos)
    antes, depois, sem_vaga, n, caps = anti_colisao(pessoas, regras, hoje)

    mudancas = []
    for p in pessoas:
        orig = sorted(p['orig'], key=lambda q: q['ini'])
        if [(q['ini'], q['dias']) for q in p['per']] == [(q['ini'], q['dias']) for q in orig]:
            continue
        for j, q in enumerate(p['per']):
            o = orig[j] if j < len(orig) else None
            if o and o['ini'] == q['ini'] and o['dias'] == q['dias']:
                continue
            fim = q['ini'] + q['dias'] - 1
            ret = fim + 1
            while dow(ret) in (0, 6) or ret in p['H']:
                ret += 1
            mudancas.append({'Colaborador': p['nome'], 'Função': p['funcao'], 'Período': j + 1, 'Dias': q['dias'],
                             'Início original': br(o['ini']) if o else '', 'Novo início': f"{DOW[dow(q['ini'])]} {br(q['ini'])}",
                             'Deslocamento (dias)': q['ini'] - o['ini'] if o else '', 'Novo fim': br(fim),
                             'Retorno': f"{DOW[dow(ret)]} {br(ret)}"})
    alterados = {p['nome']: p for p in pessoas if any(m['Colaborador'] == p['nome'] for m in mudancas)}

    out = Path(a.saida)
    out.mkdir(exist_ok=True)
    stamp = dd(hoje).isoformat()
    csv_path = out / f'alteracoes_{stamp}.csv'
    with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=list(mudancas[0].keys()) if mudancas else ['Colaborador'], delimiter=';')
        w.writeheader()
        w.writerows(mudancas)

    xlsx_path, linhas, nao_achados, faltando = None, 0, [], []
    if mudancas and not a.simular:
        xlsx_path = out / f'{src.stem}_sem_colisao_{stamp}.xlsx'
        linhas, nao_achados, faltando = gravar(src, xlsx_path, alterados)

    sem = sorted({it['p']['nome'] for it, _ in sem_vaga})
    resumo = {
        'planilha': str(src), 'referencia': br(hoje), 'regras': regras,
        'colaboradores': len(pessoas),
        'choques_antes': antes['excesso'], 'dias_com_choque_antes': antes['dias'],
        'choques_depois': depois['excesso'], 'dias_com_choque_depois': depois['dias'],
        'periodos_reprogramados': len(mudancas), 'colaboradores_reprogramados': len(alterados),
        'sem_vaga': [{'nome': it['p']['nome'], 'inicio_original': br(it['o']), 'dias': it['dias'], 'motivo': mot} for it, mot in sem_vaga],
        'limites': {k[2:] if k != 'GESTAO' else 'GESTÃO (todos juntos)': {'pessoas': n.get(k), 'limite': v,
                    'pico_antes': antes['pico'].get(k, 0), 'pico_depois': depois['pico'].get(k, 0)} for k, v in caps.items() if k != 'GESTAO' or v > 0},
        'avisos': avisos, 'csv': str(csv_path), 'planilha_ajustada': str(xlsx_path) if xlsx_path else None,
        'celulas_linhas_gravadas': linhas, 'nao_encontrados_na_planilha': nao_achados, 'celulas_faltando': faltando,
    }
    (out / f'resumo_{stamp}.json').write_text(json.dumps(resumo, ensure_ascii=False, indent=2), encoding='utf-8')

    print(f'Choques (pessoa-dias acima do limite): {antes["excesso"]} -> {depois["excesso"]}'
          f'  | dias com choque: {antes["dias"]} -> {depois["dias"]}')
    print(f'Períodos reprogramados: {len(mudancas)} de {len(alterados)} colaborador(es)')
    print(f'Sem vaga: {len(sem)}' + (f' ({", ".join(sem)})' if sem else ''))
    for av in avisos:
        print('Aviso:', av)
    print('Relatório:', csv_path)
    print('Resumo:', out / f'resumo_{stamp}.json')
    if xlsx_path:
        print(f'Planilha ajustada: {xlsx_path} ({linhas} linha(s) gravadas)')
        if nao_achados:
            print('Não encontrados na planilha:', ', '.join(nao_achados))
        for fz in faltando:
            print('Atenção:', fz)
    elif a.simular:
        print('Modo simulação: planilha ajustada não foi criada.')


if __name__ == '__main__':
    main()
